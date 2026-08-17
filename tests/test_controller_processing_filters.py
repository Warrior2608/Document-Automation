from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook

from document_automation_studio.core.controller import ApplicationController
from document_automation_studio.models.config_models import ProcessingSettings


def test_controller_skips_disabled_excel_processing(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()

    excel_file = input_dir / "workbook.xlsx"
    workbook = Workbook()
    workbook.save(excel_file)

    controller = ApplicationController(config_path=tmp_path / "config.json")
    controller.config.processing.input_folder = str(input_dir)
    controller.config.processing.output_folder = str(output_dir)
    controller.config.processing.process_excel_files = False
    controller.update_processing_settings(controller.config.processing)

    controller.run_batch()

    assert not (output_dir / "workbook.xlsx").exists()


def test_controller_skips_disabled_word_processing(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()

    word_file = input_dir / "document.docx"
    doc = Document()
    doc.add_paragraph("Hello world")
    doc.save(word_file)

    controller = ApplicationController(config_path=tmp_path / "config.json")
    controller.config.processing.input_folder = str(input_dir)
    controller.config.processing.output_folder = str(output_dir)
    controller.config.processing.process_word_files = False
    controller.update_processing_settings(controller.config.processing)

    controller.run_batch()

    assert not (output_dir / "document.docx").exists()
