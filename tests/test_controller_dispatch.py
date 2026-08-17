from __future__ import annotations

from pathlib import Path

from docx import Document

from document_automation_studio.core.controller import ApplicationController


def test_controller_dispatches_word_and_excel(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()

    word_file = input_dir / "document.docx"
    doc = Document()
    doc.add_paragraph("Company name: {{Company}}")
    doc.save(word_file)

    excel_file = input_dir / "workbook.xlsx"
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Hello"
    sheet["B1"] = "{{Company}}"
    workbook.save(excel_file)

    controller = ApplicationController(config_path=tmp_path / "config.json")
    controller.config.processing.input_folder = str(input_dir)
    controller.config.processing.output_folder = str(output_dir)
    controller.config.processing.preserve_folder_structure = False
    controller.update_processing_settings(controller.config.processing)

    controller.run_batch()

    assert (output_dir / "document.docx").exists()
    assert (output_dir / "workbook.xlsx").exists()
